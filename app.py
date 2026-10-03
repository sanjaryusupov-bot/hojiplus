# -*- coding: utf-8 -*-
"""Учёт закупа и продаж — Streamlit + Google Sheets."""
import hmac
import io
import zipfile
from datetime import date, datetime

import altair as alt
import gspread
import pandas as pd
import streamlit as st

st.set_page_config(
    page_title="Учёт закупа и продаж",
    page_icon="📦",
    layout="wide",
    initial_sidebar_state="expanded",
)

# ───────────────────────── Настройки ─────────────────────────
FIRST_ROW = 5      # первая строка с данными (4-я — заголовки)
READ_LAST = 3000   # до какой строки читаем
# Таблица зашита в приложение. В secrets можно переопределить через sheet_id.
DEFAULT_SHEET_ID = "1XDR1mYEoMIKLRdCwQkryaHEBoune5bSInP6ZrKu4pj8"
TRADE = ("Закуп", "Продажи")
SCHEMA = {
    "Закуп":   ["Дата", "Фирма", "Товар", "Цена", "Кол-во", "Сумма", "Комментарий"],
    "Продажи": ["Дата", "Фирма", "Товар", "Цена", "Кол-во", "Сумма", "Комментарий"],
    "Расходы": ["Дата", "Статья", "Сумма", "Комментарий"],
    "Касса":   ["Дата", "Операция", "Сумма", "Комментарий"],
}
NUM_COLS = {"Цена", "Кол-во", "Сумма"}
CASH_OPS = ["Взял из кассы", "Вернул в кассу"]
COLORS = {"Закуп": "#6366F1", "Продажи": "#10B981", "Расходы": "#F59E0B"}


def secret(key, default=None):
    try:
        return st.secrets[key]
    except Exception:
        return default


CURRENCY = str(secret("currency", "") or "")

# ───────────────────────── Дизайн ─────────────────────────
# ───────────────────────── Тема: день / ночь ─────────────────────────
def native_theme():
    """Какая тема сейчас у самого Streamlit (от неё зависят встроенные виджеты)."""
    try:
        t = st.context.theme.type
        return t if t in ("light", "dark") else "light"
    except Exception:  # noqa: BLE001
        return "light"


NATIVE = native_theme()
if "dark_mode" not in st.session_state:
    qp = st.query_params.get("theme")
    st.session_state["dark_mode"] = (qp == "dark") if qp in ("dark", "light") else (NATIVE == "dark")
DARK = bool(st.session_state["dark_mode"])
# Палитра всегда совпадает с родной темой Streamlit (чтобы все виджеты читались),
# а если выбран другой режим — вся страница целиком инвертируется.
INVERT = DARK != (NATIVE == "dark")

PALETTES = {
    "light": {"bg": "#F4F6FB", "card": "#FFFFFF", "sidebar": "#FFFFFF", "text": "#1F2937", "strong": "#111827",
              "muted": "#6B7280", "faint": "#9CA3AF", "border": "#E5E7EB", "tab-bg": "#EEF2FF",
              "tab-fg": "#4F46E5", "shadow": "rgba(16,24,40,.07)"},
    "dark": {"bg": "#0E1117", "card": "#1A1F2B", "sidebar": "#151A24", "text": "#E5E7EB", "strong": "#F9FAFB",
             "muted": "#9CA3AF", "faint": "#6B7280", "border": "#2A3040", "tab-bg": "#272B4D",
             "tab-fg": "#A5B4FC", "shadow": "rgba(0,0,0,.35)"},
}
_root = ":root{" + ";".join(f"--{k}:{v}" for k, v in PALETTES[NATIVE].items()) + "}"
_flip = ("html{filter:invert(1) hue-rotate(180deg);}"
         ".hero,.kpi.big{filter:invert(1) hue-rotate(180deg);}") if INVERT else ""

# ───────────────────────── Дизайн ─────────────────────────
st.markdown(
    "<style>" + _root + _flip + """
footer {visibility: hidden;}
.stApp {background: var(--bg);}
.block-container {padding-top: 1.4rem; max-width: 1280px;}
section[data-testid="stSidebar"] {background: var(--sidebar); border-right:1px solid var(--border);}

.hero {background: linear-gradient(120deg,#4F46E5 0%,#7C3AED 58%,#EC4899 125%);
  color:#fff; border-radius:20px; padding:22px 28px; margin-bottom:18px;
  display:flex; justify-content:space-between; align-items:center; gap:12px; flex-wrap:wrap;
  box-shadow:0 10px 30px rgba(79,70,229,.25);}
.hero-title {font-size:1.75rem; font-weight:800; letter-spacing:-.02em; line-height:1.2; color:#fff;}
.hero-sub {opacity:.92; margin-top:4px; font-size:.95rem; color:#fff;}
.hero-badge {background:rgba(255,255,255,.18); padding:6px 14px; border-radius:999px; font-size:.8rem; color:#fff;}

.kpi {background: var(--card); border-radius:16px; padding:16px 20px; margin-bottom:14px;
  border-left:5px solid var(--c,#6366F1);
  box-shadow:0 1px 3px var(--shadow), 0 6px 18px var(--shadow);}
.kpi-label {font-size:.74rem; color: var(--muted); font-weight:700; text-transform:uppercase; letter-spacing:.05em;}
.kpi-value {font-size:1.6rem; font-weight:800; color: var(--strong); margin-top:4px; line-height:1.15;}
.kpi-sub {font-size:.78rem; color: var(--faint); margin-top:6px;}
.kpi.indigo {--c:#6366F1;} .kpi.green {--c:#10B981;} .kpi.amber {--c:#F59E0B;}
.kpi.rose {--c:#F43F5E;} .kpi.cyan {--c:#06B6D4;} .kpi.slate {--c:#64748B;}
.kpi.big {border-left:none; background:linear-gradient(135deg,#4F46E5,#7C3AED);
  box-shadow:0 10px 26px rgba(79,70,229,.30);}
.kpi.big.neg {background:linear-gradient(135deg,#E11D48,#F97316); box-shadow:0 10px 26px rgba(225,29,72,.30);}
.kpi.big .kpi-label {color:rgba(255,255,255,.85);}
.kpi.big .kpi-value {color:#fff; font-size:2.15rem;}
.kpi.big .kpi-sub {color:rgba(255,255,255,.8);}

[class*="st-key-card_"] {background: var(--card); border-radius:16px; padding:18px 22px 10px; margin-bottom:16px;
  box-shadow:0 1px 3px var(--shadow), 0 6px 18px var(--shadow);}

.stTabs [data-baseweb="tab-list"] {gap:6px; background: var(--card); padding:6px; border-radius:14px;
  box-shadow:0 1px 3px var(--shadow);}
.stTabs [data-baseweb="tab"] {border-radius:10px; padding:8px 16px; font-weight:600; height:auto;}
.stTabs [aria-selected="true"] {background: var(--tab-bg); color: var(--tab-fg);}
.stTabs [data-baseweb="tab-highlight"], .stTabs [data-baseweb="tab-border"] {display:none;}

.stButton > button, .stFormSubmitButton > button, .stDownloadButton > button {
  border-radius:12px; font-weight:700; padding:.5rem 1.3rem;}
h5 {margin-top:0 !important;}
.section-title {font-size:1.05rem; font-weight:700; color: var(--strong); margin:6px 0 10px;}
</style>""",
    unsafe_allow_html=True,
)


def on_theme():
    st.query_params["theme"] = "dark" if st.session_state["dark_mode"] else "light"


# ───────────────────────── Вспомогательное ─────────────────────────
def money(x, dec=0):
    try:
        s = f"{float(x):,.{dec}f}".replace(",", "\u00a0")
    except Exception:
        s = "0"
    return f"{s} {CURRENCY}".strip()


def num(x):
    return f"{float(x):,.0f}".replace(",", "\u00a0")


def kpi(label, value, sub="", tone="indigo", big=False, neg=False):
    cls = f"kpi {tone}" + (" big" if big else "") + (" neg" if neg else "")
    return (f'<div class="{cls}"><div class="kpi-label">{label}</div>'
            f'<div class="kpi-value">{value}</div><div class="kpi-sub">{sub}</div></div>')


def safe(text):
    """Чтобы Google Таблица не принимала текст за формулу."""
    text = " ".join(str(text or "").split())
    return "'" + text if text[:1] in ("=", "+", "-", "@") else text


def col_letter(n):
    return chr(64 + n)


# ───────────────────────── Подключение к Google Sheets ─────────────────────────
@st.cache_resource(show_spinner=False)
def get_spreadsheet():
    creds = dict(st.secrets["gcp_service_account"])
    gc = gspread.service_account_from_dict(creds)
    sid = str(secret("sheet_id", DEFAULT_SHEET_ID)).strip()
    return gc.open_by_url(sid) if sid.startswith("http") else gc.open_by_key(sid)


def parse_dates(s):
    if s.empty:
        return pd.Series(pd.to_datetime([]), index=s.index)
    n = pd.to_numeric(s, errors="coerce")
    d1 = pd.to_datetime(n, unit="D", origin="1899-12-30", errors="coerce")
    txt = s.where(n.isna()).astype(str).replace({"nan": None, "None": None, "": None})
    d2 = pd.to_datetime(txt, dayfirst=True, errors="coerce", format="mixed")
    return d1.fillna(d2).dt.normalize()


def to_num(s):
    if s.empty:
        return s.astype(float)
    t = (s.astype(str).str.replace("\u00a0", "", regex=False)
         .str.replace(" ", "", regex=False).str.replace(",", ".", regex=False))
    return pd.to_numeric(t, errors="coerce")


def to_frame(name, rows):
    cols = SCHEMA[name]
    n = len(cols)
    rows = [(list(r) + [""] * n)[:n] for r in rows]
    raw = pd.DataFrame(rows, columns=cols)
    df = pd.DataFrame({"_row": range(FIRST_ROW, FIRST_ROW + len(raw))})
    for c in cols:
        if c == "Дата":
            df[c] = parse_dates(raw[c])
        elif c in NUM_COLS:
            df[c] = to_num(raw[c])
        else:
            df[c] = raw[c].astype(str).str.strip().replace({"nan": "", "None": ""})
    if name in TRADE:
        df["Сумма"] = df["Цена"].fillna(0) * df["Кол-во"].fillna(0)
        has = df["Цена"].notna() | df["Кол-во"].notna()
        df["Товар"] = df["Товар"].replace("", "—")
    else:
        has = df["Сумма"].notna()
    nodate = int((has & df["Дата"].isna()).sum())
    df = df[has & df["Дата"].notna()].reset_index(drop=True)
    for c in NUM_COLS & set(df.columns):
        df[c] = df[c].astype(float)
    return df, nodate


@st.cache_data(ttl=120, show_spinner="Загружаю данные из таблицы…")
def load_data():
    sh = get_spreadsheet()
    ranges = [f"'{n}'!A{FIRST_ROW}:{col_letter(len(c))}{READ_LAST}" for n, c in SCHEMA.items()]
    res = sh.values_batch_get(
        ranges, params={"valueRenderOption": "UNFORMATTED_VALUE", "dateTimeRenderOption": "SERIAL_NUMBER"})
    out, nodate = {}, 0
    for name, vr in zip(SCHEMA, res["valueRanges"]):
        out[name], k = to_frame(name, vr.get("values", []))
        nodate += k
    out["_nodate"] = nodate
    out["_loaded_at"] = datetime.now()
    return out


def refresh():
    st.cache_data.clear()


# ───────────────────────── Запись в таблицу ─────────────────────────
def write_row(sheet, values):
    """Пишет запись в первую свободную строку. values — без колонки «Сумма» для Закуп/Продажи."""
    ws = get_spreadsheet().worksheet(sheet)
    key_end = "E" if sheet in TRADE else "C"
    rows = ws.get(f"A{FIRST_ROW}:{key_end}{READ_LAST}", value_render_option="UNFORMATTED_VALUE")
    r = FIRST_ROW + len(rows)
    for i, row in enumerate(rows):
        if all(str(c).strip() == "" for c in row):
            r = FIRST_ROW + i
            break
    if sheet in TRADE:
        d, firm, prod, price, qty, comment = values
        full = [d, firm, prod, price, qty, f'=IF(OR(D{r}="",E{r}=""),"",D{r}*E{r})', comment]
        rng = f"A{r}:G{r}"
    else:
        full = values
        rng = f"A{r}:D{r}"
    ws.batch_update([{"range": rng, "values": [full]}], value_input_option="USER_ENTERED")
    return r


def set_flash(sheet, kind, msg):
    st.session_state[f"flash_{sheet}"] = (kind, msg)


def names_by_freq(*series):
    s = pd.concat(series) if series else pd.Series(dtype=str)
    s = s[(s != "") & (s != "—")]
    return s.value_counts().index.tolist()


def canon(value, options):
    v = " ".join(str(value or "").split())
    for o in options:
        if o.lower() == v.lower():
            return o
    return v


def save_entry(sheet):
    ss = st.session_state
    k = lambda n: f"{sheet}__{n}"
    data = load_data()
    d = ss.get(k("date")) or date.today()
    comment = safe(ss.get(k("comment")))
    try:
        if sheet in TRADE:
            firm = canon(ss.get(k("firm")), names_by_freq(data[sheet]["Фирма"]))
            prod = canon(ss.get(k("prod")), names_by_freq(data["Закуп"]["Товар"], data["Продажи"]["Товар"])) or "—"
            price, qty = ss.get(k("price")), ss.get(k("qty"))
            if not firm:
                return set_flash(sheet, "err", "Укажи фирму — выбери из списка или впиши новую.")
            if price is None or price <= 0:
                return set_flash(sheet, "err", "Укажи цену за 1 шт.")
            if qty is None or qty <= 0:
                return set_flash(sheet, "err", "Укажи количество.")
            write_row(sheet, [d.isoformat(), safe(firm), safe(prod), price, int(qty), comment])
            msg = f"{sheet}: {firm} · {prod} · {money(price * qty)}"
            for n in ("prod", "price", "qty", "comment"):
                ss[k(n)] = "" if n == "comment" else None
        else:
            amount = ss.get(k("sum"))
            if amount is None or amount <= 0:
                return set_flash(sheet, "err", "Укажи сумму.")
            if sheet == "Расходы":
                name = canon(ss.get(k("name")), names_by_freq(data["Расходы"]["Статья"]))
                if not name:
                    return set_flash(sheet, "err", "Укажи статью расхода — выбери из списка или впиши новую.")
                ss[k("name")] = None
            else:
                name = ss.get(k("name")) or CASH_OPS[0]
            write_row(sheet, [d.isoformat(), safe(name), amount, comment])
            msg = f"{sheet}: {name} · {money(amount)}"
            ss[k("sum")] = None
            ss[k("comment")] = ""
    except Exception as e:  # noqa: BLE001
        return set_flash(sheet, "err", f"Не удалось записать в таблицу: {e}")
    st.cache_data.clear()
    set_flash(sheet, "ok", f"Сохранено · {msg}")


def undo_last(sheet):
    try:
        df = load_data()[sheet]
        if df.empty:
            return
        r = int(df["_row"].max())
        ws = get_spreadsheet().worksheet(sheet)
        ws.batch_clear([f"A{r}:E{r}", f"G{r}"] if sheet in TRADE else [f"A{r}:D{r}"])
        st.cache_data.clear()
        set_flash(sheet, "ok", f"Строка {r} удалена из «{sheet}»")
    except Exception as e:  # noqa: BLE001
        set_flash(sheet, "err", f"Не удалось удалить: {e}")


def show_flash(sheet):
    f = st.session_state.pop(f"flash_{sheet}", None)
    if not f:
        return
    if f[0] == "ok":
        st.toast(f[1], icon="✅")
    else:
        st.error(f[1], icon="⚠️")


# ───────────────────────── Расчёты ─────────────────────────
def between(df, start, end):
    return df[(df["Дата"] >= start) & (df["Дата"] <= end)]


def upto(df, end):
    return df[df["Дата"] <= end]


def cash_net(df):
    return (df.loc[df["Операция"] == CASH_OPS[0], "Сумма"].sum()
            - df.loc[df["Операция"] == CASH_OPS[1], "Сумма"].sum())


def avg_cost(pur):
    g = pur.groupby("Товар").agg(s=("Сумма", "sum"), q=("Кол-во", "sum"))
    return (g["s"] / g["q"].where(g["q"] > 0)).fillna(0)


def compute(data, start, end):
    pur_p, sal_p = between(data["Закуп"], start, end), between(data["Продажи"], start, end)
    exp_p, cash_p = between(data["Расходы"], start, end), between(data["Касса"], start, end)
    pur_c, sal_c = upto(data["Закуп"], end), upto(data["Продажи"], end)
    exp_c, cash_c = upto(data["Расходы"], end), upto(data["Касса"], end)

    cost = avg_cost(pur_c)
    cogs = float((sal_p["Кол-во"] * sal_p["Товар"].map(cost).fillna(0)).sum())
    revenue = float(sal_p["Сумма"].sum())
    gross = revenue - cogs
    expenses = float(exp_p["Сумма"].sum())
    stock_qty = float(pur_c["Кол-во"].sum() - sal_c["Кол-во"].sum())
    total_cost = float(pur_c["Сумма"].sum())
    stock_val = stock_qty * (total_cost / pur_c["Кол-во"].sum()) if pur_c["Кол-во"].sum() > 0 else 0.0
    return dict(
        pur_p=pur_p, sal_p=sal_p, exp_p=exp_p, cash_p=cash_p,
        purchase=float(pur_p["Сумма"].sum()), revenue=revenue, expenses=expenses,
        cogs=cogs, gross=gross, net=gross - expenses,
        margin=(gross / revenue * 100) if revenue > 0 else 0.0,
        balance=cash_net(cash_c) + float(sal_c["Сумма"].sum()) - total_cost - float(exp_c["Сумма"].sum()),
        budget_total=float(cash_net(cash_c)), budget_period=float(cash_net(cash_p)),
        stock_qty=stock_qty, stock_val=stock_val,
    )


def products_table(data, start, end):
    pur_c, sal_c = upto(data["Закуп"], end), upto(data["Продажи"], end)
    sal_p = between(data["Продажи"], start, end)
    names = sorted(set(pur_c["Товар"]) | set(sal_c["Товар"]))
    cols = ["Товар", "Закуплено, шт", "Продано всего, шт", "Остаток, шт", "Ср. цена закупа",
            "Стоимость остатка", "Продано за период, шт", "Выручка", "Себестоимость", "Прибыль", "Маржа, %"]
    if not names:
        return pd.DataFrame(columns=cols)
    t = pd.DataFrame(index=pd.Index(names, name="Товар"))
    t["Закуплено, шт"] = pur_c.groupby("Товар")["Кол-во"].sum()
    t["Продано всего, шт"] = sal_c.groupby("Товар")["Кол-во"].sum()
    t = t.fillna(0)
    t["Остаток, шт"] = t["Закуплено, шт"] - t["Продано всего, шт"]
    t["Ср. цена закупа"] = (pur_c.groupby("Товар")["Сумма"].sum() / t["Закуплено, шт"].where(t["Закуплено, шт"] > 0)).fillna(0)
    t["Стоимость остатка"] = t["Остаток, шт"] * t["Ср. цена закупа"]
    t["Продано за период, шт"] = sal_p.groupby("Товар")["Кол-во"].sum()
    t["Выручка"] = sal_p.groupby("Товар")["Сумма"].sum()
    t = t.fillna(0)
    t["Себестоимость"] = t["Продано за период, шт"] * t["Ср. цена закупа"]
    t["Прибыль"] = t["Выручка"] - t["Себестоимость"]
    t["Маржа, %"] = (t["Прибыль"] / t["Выручка"].where(t["Выручка"] > 0) * 100).fillna(0)
    return t.reset_index()[cols]


# ───────────────────────── Экспорт ─────────────────────────
def export_df(df):
    d = df.drop(columns="_row", errors="ignore").sort_values("Дата").copy()
    d["Дата"] = d["Дата"].dt.date
    return d


@st.cache_data(show_spinner=False)
def build_xlsx(frames, summary, products):
    from openpyxl.styles import Alignment, Font, PatternFill
    bio = io.BytesIO()
    with pd.ExcelWriter(bio, engine="openpyxl") as xw:
        summary.to_excel(xw, sheet_name="Сводка", index=False)
        for name, df in frames.items():
            export_df(df).to_excel(xw, sheet_name=name, index=False)
        products.to_excel(xw, sheet_name="Товары", index=False)
        for ws in xw.book.worksheets:
            for c in ws[1]:
                c.font = Font(name="Arial", bold=True, color="FFFFFF", size=10)
                c.fill = PatternFill("solid", fgColor="4F46E5")
                c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            for col in ws.columns:
                w = max((len(str(c.value)) if c.value is not None else 0) for c in col)
                ws.column_dimensions[col[0].column_letter].width = min(max(w + 3, 12), 48)
            for row in ws.iter_rows(min_row=2):
                for c in row:
                    c.font = Font(name="Arial", size=10)
                    if isinstance(c.value, date):
                        c.number_format = "DD.MM.YYYY"
                    elif isinstance(c.value, float):
                        c.number_format = "#,##0" if float(c.value).is_integer() else "#,##0.00"
            ws.freeze_panes = "A2"
    return bio.getvalue()


@st.cache_data(show_spinner=False)
def build_csv_zip(frames):
    bio = io.BytesIO()
    with zipfile.ZipFile(bio, "w", zipfile.ZIP_DEFLATED) as z:
        for name, df in frames.items():
            z.writestr(f"{name}.csv", export_df(df).to_csv(index=False, sep=";").encode("utf-8-sig"))
    return bio.getvalue()


# ───────────────────────── Графики ─────────────────────────
def dynamics_chart(m, start, end):
    daily = (end - start).days <= 62
    parts = []
    for name, df in (("Закуп", m["pur_p"]), ("Продажи", m["sal_p"]), ("Расходы", m["exp_p"])):
        if df.empty:
            continue
        per = df["Дата"] if daily else df["Дата"].dt.to_period("M").dt.to_timestamp()
        g = df.groupby(per)["Сумма"].sum().reset_index()
        g.columns = ["_p", "Сумма"]
        g["Тип"] = name
        parts.append(g)
    if not parts:
        return None
    long = pd.concat(parts, ignore_index=True)
    long["Период"] = long["_p"].dt.strftime("%d.%m" if daily else "%m.%Y")
    order = long.drop_duplicates("_p").sort_values("_p")["Период"].tolist()
    return (
        alt.Chart(long[["Период", "Тип", "Сумма"]])
        .mark_bar(cornerRadiusTopLeft=5, cornerRadiusTopRight=5)
        .encode(
            x=alt.X("Период:N", sort=order, title=None, axis=alt.Axis(labelAngle=0)),
            xOffset="Тип:N",
            y=alt.Y("Сумма:Q", title=None, axis=alt.Axis(format="~s")),
            color=alt.Color("Тип:N", legend=alt.Legend(orient="top", title=None),
                            scale=alt.Scale(domain=list(COLORS), range=list(COLORS.values()))),
            tooltip=["Период", "Тип", alt.Tooltip("Сумма:Q", format=",.0f")],
        )
        .properties(height=300)
    )


def hbar(df, label, color):
    return (
        alt.Chart(df)
        .mark_bar(cornerRadiusEnd=5, color=color)
        .encode(
            y=alt.Y(f"{label}:N", sort="-x", title=None),
            x=alt.X("Сумма:Q", title=None, axis=alt.Axis(format="~s")),
            tooltip=[label, alt.Tooltip("Сумма:Q", format=",.0f")],
        )
        .properties(height=max(110, 38 * len(df)))
    )


def top_chart(df, label, color):
    if df.empty:
        st.caption("Нет данных за период")
        return
    g = df.groupby(label, as_index=False)["Сумма"].sum().sort_values("Сумма", ascending=False).head(7)
    st.altair_chart(hbar(g, label, color))


# ───────────────────────── Формы и таблицы ─────────────────────────
def trade_form(sheet, data):
    k = lambda n: f"{sheet}__{n}"
    firms = names_by_freq(data[sheet]["Фирма"])
    prods = names_by_freq(data["Закуп"]["Товар"], data["Продажи"]["Товар"])
    firm_label = "Фирма (поставщик)" if sheet == "Закуп" else "Фирма (покупатель)"
    with st.container(key=f"card_{sheet}_form"):
        st.markdown("##### ➕ Новая запись")
        with st.form(f"form_{sheet}", border=False):
            c1, c2, c3 = st.columns([1, 1.5, 1.5])
            c1.date_input("Дата", value=date.today(), format="DD.MM.YYYY", key=k("date"))
            c2.selectbox(firm_label, firms, index=None, accept_new_options=True,
                         placeholder="Выбери из списка или впиши новую", key=k("firm"))
            c3.selectbox("Товар", prods, index=None, accept_new_options=True,
                         placeholder="Выбери из списка или впиши новый", key=k("prod"))
            c4, c5, c6 = st.columns([1, 1, 2])
            c4.number_input("Цена за 1 шт", min_value=0.0, value=None, step=1.0, placeholder="0", key=k("price"))
            c5.number_input("Всего шт", min_value=0, value=None, step=1, placeholder="0", key=k("qty"))
            c6.text_input("Комментарий", placeholder="необязательно", key=k("comment"))
            st.form_submit_button("💾  Сохранить", type="primary", on_click=save_entry, args=(sheet,))


def simple_form(sheet, data):
    k = lambda n: f"{sheet}__{n}"
    with st.container(key=f"card_{sheet}_form"):
        st.markdown("##### ➕ Новая запись")
        with st.form(f"form_{sheet}", border=False):
            c1, c2, c3 = st.columns([1, 1.6, 1.2])
            c1.date_input("Дата", value=date.today(), format="DD.MM.YYYY", key=k("date"))
            if sheet == "Расходы":
                c2.selectbox("Название расхода", names_by_freq(data["Расходы"]["Статья"]), index=None,
                             accept_new_options=True, placeholder="Выбери из списка или впиши новый", key=k("name"))
            else:
                c2.selectbox("Операция", CASH_OPS, index=0, key=k("name"))
            c3.number_input("Сумма", min_value=0.0, value=None, step=1.0, placeholder="0", key=k("sum"))
            st.text_input("Комментарий", placeholder="необязательно", key=k("comment"))
            st.form_submit_button("💾  Сохранить", type="primary", on_click=save_entry, args=(sheet,))


def show_table(df, sheet):
    if df.empty:
        st.info("За выбранный период записей нет.", icon="🗓️")
        return
    q = st.text_input("Поиск", key=f"{sheet}__q", placeholder="🔎  Поиск по фирме, товару, комментарию…",
                      label_visibility="collapsed")
    view = df
    if q:
        ql = q.lower()
        mask = pd.Series(False, index=df.index)
        for c in df.columns:
            if df[c].dtype == object:
                mask |= df[c].astype(str).str.lower().str.contains(ql, regex=False)
        view = df[mask]
    view = view.sort_values(["Дата", "_row"], ascending=False).drop(columns="_row")
    cfg = {
        "Дата": st.column_config.DateColumn("Дата", format="DD.MM.YYYY"),
        "Цена": st.column_config.NumberColumn("Цена за 1 шт", format="localized"),
        "Кол-во": st.column_config.NumberColumn("Всего шт", format="%d"),
        "Сумма": st.column_config.NumberColumn("Сумма", format="localized"),
    }
    st.dataframe(view, hide_index=True, column_config=cfg, height=min(560, 44 + 35 * len(view)))
    st.caption(f"Записей: {len(view)} · Итого: {money(view['Сумма'].sum())}")


def undo_block(sheet, data):
    df = data[sheet]
    if df.empty:
        return
    last = df.loc[df["_row"].idxmax()]
    who = last["Фирма"] if "Фирма" in df else (last["Статья"] if "Статья" in df else last["Операция"])
    with st.popover("↩️ Отменить последнюю запись"):
        st.write(f"Будет очищена строка **{int(last['_row'])}** в таблице:")
        st.caption(f"{last['Дата']:%d.%m.%Y} · {who} · {money(last['Сумма'])}")
        st.button("Да, удалить", key=f"{sheet}__undo", on_click=undo_last, args=(sheet,), type="primary")


def trade_tab(sheet, data, start, end):
    show_flash(sheet)
    trade_form(sheet, data)
    df = between(data[sheet], start, end)
    qty, total = df["Кол-во"].sum(), df["Сумма"].sum()
    tone = "indigo" if sheet == "Закуп" else "green"
    a, b, c = st.columns(3)
    a.markdown(kpi(f"{sheet} за период", money(total), f"{len(df)} записей", tone), unsafe_allow_html=True)
    b.markdown(kpi("Всего штук", num(qty), "за период", "cyan"), unsafe_allow_html=True)
    c.markdown(kpi("Средняя цена за 1 шт", money(total / qty if qty else 0, 2), "за период", "slate"),
               unsafe_allow_html=True)
    show_table(df, sheet)
    undo_block(sheet, data)


def simple_tab(sheet, data, start, end):
    show_flash(sheet)
    simple_form(sheet, data)
    df = between(data[sheet], start, end)
    if sheet == "Расходы":
        top = df.groupby("Статья")["Сумма"].sum().sort_values(ascending=False)
        a, b, c = st.columns(3)
        a.markdown(kpi("Расходы за период", money(df["Сумма"].sum()), f"{len(df)} записей", "amber"), unsafe_allow_html=True)
        b.markdown(kpi("Самая большая статья", top.index[0] if len(top) else "—",
                       money(top.iloc[0]) if len(top) else "", "rose"), unsafe_allow_html=True)
        c.markdown(kpi("Средний расход", money(df["Сумма"].mean() if len(df) else 0), "на одну запись", "slate"),
                   unsafe_allow_html=True)
    else:
        took = df.loc[df["Операция"] == CASH_OPS[0], "Сумма"].sum()
        back = df.loc[df["Операция"] == CASH_OPS[1], "Сумма"].sum()
        a, b, c = st.columns(3)
        a.markdown(kpi("Взято за период", money(took), "", "indigo"), unsafe_allow_html=True)
        b.markdown(kpi("Возвращено за период", money(back), "", "green"), unsafe_allow_html=True)
        c.markdown(kpi("Бюджет всего (до конца периода)", money(st.session_state["_budget_total"]),
                       "взято − возвращено", "slate"), unsafe_allow_html=True)
    show_table(df, sheet)
    undo_block(sheet, data)


def overview(m, start, end):
    balance_sub = "бюджет + продажи − закуп − расходы · на конец периода"
    c = st.columns([1.35, 1, 1, 1])
    c[0].markdown(kpi("Основной баланс", money(m["balance"]), balance_sub, big=True, neg=m["balance"] < 0),
                  unsafe_allow_html=True)
    c[1].markdown(kpi("Общий закуп", money(m["purchase"]),
                      f"{num(m['pur_p']['Кол-во'].sum())} шт · {len(m['pur_p'])} записей", "indigo"),
                  unsafe_allow_html=True)
    c[2].markdown(kpi("Продажи", money(m["revenue"]),
                      f"{num(m['sal_p']['Кол-во'].sum())} шт · {len(m['sal_p'])} записей", "green"),
                  unsafe_allow_html=True)
    c[3].markdown(kpi("Расходы", money(m["expenses"]), f"{len(m['exp_p'])} записей", "amber"),
                  unsafe_allow_html=True)
    c = st.columns([1.35, 1, 1, 1])
    c[0].markdown(kpi("Чистая прибыль", money(m["net"]), "продажи − себестоимость − расходы", "green" if m["net"] >= 0 else "rose"),
                  unsafe_allow_html=True)
    c[1].markdown(kpi("Валовая прибыль", money(m["gross"]), f"себестоимость: {money(m['cogs'])}", "cyan"),
                  unsafe_allow_html=True)
    c[2].markdown(kpi("Маржа с продаж", f"{m['margin']:.1f} %", "валовая прибыль / продажи", "indigo"),
                  unsafe_allow_html=True)
    c[3].markdown(kpi("Остаток на складе", f"{num(m['stock_qty'])} шт", f"на сумму {money(m['stock_val'])}", "slate"),
                  unsafe_allow_html=True)

    left, right = st.columns([2, 1])
    with left.container(key="card_dyn"):
        st.markdown('<div class="section-title">Закуп, продажи и расходы по времени</div>', unsafe_allow_html=True)
        ch = dynamics_chart(m, start, end)
        if ch is None:
            st.info("За выбранный период данных нет.", icon="🗓️")
        else:
            st.altair_chart(ch)
    with right.container(key="card_exp"):
        st.markdown('<div class="section-title">Расходы по статьям</div>', unsafe_allow_html=True)
        top_chart(m["exp_p"], "Статья", COLORS["Расходы"])

    a, b = st.columns(2)
    with a.container(key="card_sup"):
        st.markdown('<div class="section-title">Топ поставщиков</div>', unsafe_allow_html=True)
        top_chart(m["pur_p"], "Фирма", COLORS["Закуп"])
    with b.container(key="card_buy"):
        st.markdown('<div class="section-title">Топ покупателей</div>', unsafe_allow_html=True)
        top_chart(m["sal_p"], "Фирма", COLORS["Продажи"])
    st.caption("Себестоимость считается по средней цене закупа каждого товара (по всем закупам до конца периода).")


def products_tab(data, start, end):
    t = products_table(data, start, end)
    if t.empty:
        st.info("Пока нет закупов и продаж.", icon="📦")
        return
    cfg = {c: st.column_config.NumberColumn(c, format="localized") for c in t.columns if c != "Товар"}
    for c in ("Закуплено, шт", "Продано всего, шт", "Остаток, шт", "Продано за период, шт"):
        cfg[c] = st.column_config.NumberColumn(c, format="%d")
    cfg["Маржа, %"] = st.column_config.NumberColumn("Маржа, %", format="%.1f%%")
    st.dataframe(t, hide_index=True, column_config=cfg, height=min(600, 44 + 35 * len(t)))
    st.caption("Остаток и средняя цена закупа — накопительно до конца периода. "
               "Продажи, прибыль и маржа — только за выбранный период.")


# ───────────────────────── Вход по паролю ─────────────────────────
def gate():
    pwd = secret("app_password")
    if not pwd or st.session_state.get("auth"):
        return
    st.markdown('<div class="hero"><div><div class="hero-title">📦 Учёт закупа и продаж</div>'
                '<div class="hero-sub">Введи пароль, чтобы продолжить</div></div></div>', unsafe_allow_html=True)
    with st.form("login"):
        p = st.text_input("Пароль", type="password")
        ok = st.form_submit_button("Войти", type="primary")
    if ok:
        if hmac.compare_digest(p.encode(), str(pwd).encode()):
            st.session_state["auth"] = True
            st.rerun()
        else:
            st.error("Неверный пароль")
    st.stop()


# ───────────────────────── Запуск ─────────────────────────
gate()

try:
    data = load_data()
except KeyError as e:
    st.error(f"В Secrets не хватает: {e}. Добавь блок [gcp_service_account] с данными из JSON-ключа.", icon="🔑")
    st.stop()
except gspread.exceptions.WorksheetNotFound as e:
    st.error(f"В таблице нет вкладки {e}. Нужны вкладки: Закуп, Продажи, Расходы, Касса.", icon="📄")
    st.stop()
except Exception as e:  # noqa: BLE001
    st.error(f"Не получилось подключиться к Google Таблице: {e}\n\n"
             "Проверь, что таблица открыта для сервисного аккаунта (доступ «Редактор») и ID верный.", icon="🔌")
    st.button("Повторить", on_click=refresh)
    st.stop()

# ── Боковая панель: период и выгрузка
today = pd.Timestamp(date.today())
with st.sidebar:
    st.markdown("### 📦 Учёт")
    choice = st.radio("Период", ["Сегодня", "7 дней", "Этот месяц", "Прошлый месяц", "Всё время", "Свой период"],
                      index=2)
    first_of_month = today.replace(day=1)
    if choice == "Сегодня":
        start, end = today, today
    elif choice == "7 дней":
        start, end = today - pd.Timedelta(days=6), today
    elif choice == "Этот месяц":
        start, end = first_of_month, today
    elif choice == "Прошлый месяц":
        end = first_of_month - pd.Timedelta(days=1)
        start = end.replace(day=1)
    elif choice == "Всё время":
        all_dates = pd.concat([data[s]["Дата"] for s in SCHEMA])
        start = all_dates.min() if len(all_dates) else today
        end = max(today, all_dates.max()) if len(all_dates) else today
    else:
        rng = st.date_input("Даты", value=(first_of_month.date(), today.date()), format="DD.MM.YYYY")
        if not isinstance(rng, (tuple, list)) or len(rng) != 2:
            st.info("Выбери вторую дату периода")
            st.stop()
        start, end = pd.Timestamp(rng[0]), pd.Timestamp(rng[1])
    st.button("🔄 Обновить данные", on_click=refresh)
    if data["_nodate"]:
        st.warning(f"В таблице {data['_nodate']} строк(и) без даты — они не учтены.", icon="⚠️")

m = compute(data, start, end)
st.session_state["_budget_total"] = m["budget_total"]
frames = {"Закуп": m["pur_p"], "Продажи": m["sal_p"], "Расходы": m["exp_p"], "Касса": m["cash_p"]}
products = products_table(data, start, end)
period_txt = f"{start:%d.%m.%Y} — {end:%d.%m.%Y}" if start != end else f"{start:%d.%m.%Y}"

summary = pd.DataFrame({
    "Показатель": ["Период", "Основной баланс (на конец периода)", "Бюджет взят из кассы (за период)", "Общий закуп",
                   "Продажи", "Расходы", "Себестоимость проданного", "Валовая прибыль", "Чистая прибыль",
                   "Маржа с продаж, %", "Остаток на складе, шт", "Стоимость остатка"],
    "Значение": [period_txt, m["balance"], m["budget_period"], m["purchase"], m["revenue"], m["expenses"],
                 m["cogs"], m["gross"], m["net"], round(m["margin"], 2), m["stock_qty"], m["stock_val"]],
})
fname = f"uchet_{start:%Y%m%d}_{end:%Y%m%d}"
with st.sidebar:
    st.divider()
    st.markdown("**📥 Выгрузка за период**")
    st.download_button("Скачать Excel (.xlsx)", build_xlsx(frames, summary, products), file_name=f"{fname}.xlsx",
                       mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    st.download_button("Скачать CSV (.zip)", build_csv_zip(frames), file_name=f"{fname}.zip",
                       mime="application/zip")

# ── Переключатель день / ночь + шапка
_sp, _tg = st.columns([5, 1.6])
_tg.toggle("🌙 Ночной режим", key="dark_mode", on_change=on_theme)
st.markdown(
    f'<div class="hero"><div><div class="hero-title">📦 Учёт закупа и продаж</div>'
    f'<div class="hero-sub">Период: {period_txt}</div></div>'
    f'<div class="hero-badge">Обновлено {data["_loaded_at"]:%H:%M:%S}</div></div>',
    unsafe_allow_html=True,
)

tabs = st.tabs(["📊 Обзор", "🛒 Закуп", "💰 Продажи", "💸 Расходы", "🏦 Касса", "📦 Товары"])
with tabs[0]:
    overview(m, start, end)
with tabs[1]:
    trade_tab("Закуп", data, start, end)
with tabs[2]:
    trade_tab("Продажи", data, start, end)
with tabs[3]:
    simple_tab("Расходы", data, start, end)
with tabs[4]:
    simple_tab("Касса", data, start, end)
with tabs[5]:
    products_tab(data, start, end)
